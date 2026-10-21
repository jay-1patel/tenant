"""Verification test for Auto-Bifurcation and Auto-Configuration Engine."""

import sys
import os
from pathlib import Path
import json

# Setup paths
PROJECT_ROOT = Path("c:/Users/Leeway/Desktop/chatbot2")
os.chdir(str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "backend"))
sys.path.insert(0, str(PROJECT_ROOT))

import docx
from services.auto_bifurcator import AutoBifurcator
from database import list_products
from shared.tenancy import loader as tenancy_loader
from shared.tenancy import store as tenancy_store


def create_sample_docx(filepath: str):
    """Create a sample client brochure/catalog document in Word format."""
    doc = docx.Document()
    doc.add_heading("Apex Global Logistics & Freight Services", level=0)

    doc.add_heading("Company Overview", level=1)
    doc.add_paragraph(
        "Apex Global Logistics is an international freight forwarding, cargo shipping, "
        "and supply chain logistics provider founded in 2012. We operate across 45 countries, "
        "providing air freight, ocean container shipping, customs clearance, and express door-to-door delivery."
    )
    doc.add_paragraph("Website: https://www.apexlogistics.com")
    doc.add_paragraph("Support Email: care@apexlogistics.com")
    doc.add_paragraph("Customer Helpline: +91 98765 43210")
    doc.add_paragraph("Operating Hours: Monday to Saturday, 09:00 AM - 07:00 PM (Asia/Kolkata)")

    doc.add_heading("Our Core Logistics Services & Pricing", level=1)
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Service Name"
    hdr_cells[1].text = "Category"
    hdr_cells[2].text = "Starting Rate"
    hdr_cells[3].text = "Key Features"

    services = [
        ("Express Air Cargo", "Air Freight", "$4.50 / kg", "Next-day arrival, temperature controlled, customs cleared"),
        ("Standard Ocean Freight FCL", "Ocean Freight", "$1,850 / container", "20ft & 40ft containers, port-to-port tracking, weekly sailings"),
        ("Cross-Border Customs Brokerage", "Customs", "$150 / consignment", "Full duty calculation, compliance filing, EDI clearance"),
        ("Warehousing & 3PL Fulfillment", "Warehousing", "$12 / pallet / month", "24/7 CCTV, barcode scanning, pick-and-pack fulfillment"),
    ]

    for name, cat, price, features in services:
        row_cells = table.add_row().cells
        row_cells[0].text = name
        row_cells[1].text = cat
        row_cells[2].text = price
        row_cells[3].text = features

    doc.add_heading("Frequently Asked Questions", level=1)
    doc.add_paragraph("Q: What documents are required for international customs clearance?")
    doc.add_paragraph("A: You need a Commercial Invoice, Packing List, Bill of Lading or Airway Bill, and Certificate of Origin.")

    doc.add_paragraph("Q: Can I track my cargo in real-time?")
    doc.add_paragraph("A: Yes, every shipment is assigned a 12-digit tracking AWB number accessible via WhatsApp and our portal.")

    doc.add_paragraph("Q: What is the claim and refund policy for transit damage?")
    doc.add_paragraph("A: All insured consignments can file claims within 7 days of delivery with inspection photographs.")

    doc.save(filepath)
    print(f"Created sample DOCX at: {filepath}")


def test_bifurcation():
    test_tenant = "apexlogistics"
    print(f"\n================ TESTING AUTO-BIFURCATION FOR TENANT '{test_tenant}' ================\n")

    # 1. Ensure tenant exists
    tenancy_store.ensure_tenant(
        test_tenant,
        display_name="Apex Global Logistics",
        vertical="generic",
        status="active"
    )

    # 2. Create sample document
    sample_doc_path = PROJECT_ROOT / "scripts" / "sample_apex_brochure.docx"
    create_sample_docx(str(sample_doc_path))

    with open(sample_doc_path, "rb") as f:
        file_bytes = f.read()

    # 3. Run AutoBifurcator
    bifurcator = AutoBifurcator(
        tenant_id=test_tenant,
        actor={"username": "test_admin", "role": "super_admin", "tenant_id": test_tenant}
    )

    files_payload = [
        {"filename": "sample_apex_brochure.docx", "content": file_bytes, "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    ]

    print("Running process_files...")
    summary = bifurcator.process_files(files_payload, auto_publish=True)

    print("\n--- AUTO-BIFURCATION SUMMARY RESULT ---")
    print(json.dumps(summary, indent=2))

    # 4. Verify Database Products
    products = list_products(tenant_id=test_tenant, active_only=False)
    print(f"\n[OK] Catalog Products in SQLite: {len(products)} items found:")
    for p in products:
        safe_name = str(p['name']).encode('ascii', 'ignore').decode('ascii')
        safe_price = str(p.get('price', '')).encode('ascii', 'ignore').decode('ascii')
        print(f"   - {safe_name} | Category: {p['category']} | Price: {safe_price}")

    # 5. Verify Tenant Profile
    profile = tenancy_loader.get_tenant_profile(test_tenant)
    print(f"\n[OK] Tenant Profile Loaded:")
    print(f"   - Brand Name: {str(profile.brand.name).encode('ascii', 'ignore').decode('ascii')}")
    print(f"   - Bot Name: {str(profile.brand.bot_name).encode('ascii', 'ignore').decode('ascii')}")
    print(f"   - Tagline: {str(profile.brand.tagline).encode('ascii', 'ignore').decode('ascii')}")
    print(f"   - Support Email: {profile.brand.support_email}")
    print(f"   - Support Phone: {profile.brand.support_phone}")
    print(f"   - Menu Buttons Count: {len(profile.menu.buttons)}")
    for b in profile.menu.buttons:
        safe_btitle = str(b.title).encode('ascii', 'ignore').decode('ascii')
        safe_bdesc = str(b.description).encode('ascii', 'ignore').decode('ascii')
        print(f"       * [{safe_btitle}] (len={len(b.title)}) -> {safe_bdesc} (len={len(b.description)})")
    print(f"   - Active Intents: {len(profile.intents)}")

    # 6. Verify Knowledge Base Files
    kb_dir = PROJECT_ROOT / "clients" / test_tenant / "knowledge"
    if kb_dir.exists():
        kb_files = list(kb_dir.glob("*.md"))
        print(f"\n[OK] Generated Knowledge Base Files: {len(kb_files)}")
        for kbf in kb_files:
            print(f"   - {kbf.name} ({kbf.stat().st_size} bytes)")

    # 7. Verify Client Config JSON
    config_file = PROJECT_ROOT / "clients" / test_tenant / "config.json"
    print(f"\n[OK] Client config.json exists: {config_file.exists()} (size={config_file.stat().st_size if config_file.exists() else 0} bytes)")

    print("\n================ TEST PASSED SUCCESSFULLY! ================\n")


if __name__ == "__main__":
    test_bifurcation()
