import os
import re
import json
import sqlite3
from typing import List, Dict
from threading import Lock

# Force PyTorch to use CPU-only mode to prevent CUDA initialization issues
os.environ["CUDA_VISIBLE_DEVICES"] = ""
os.environ["TORCH_DEVICE"] = "cpu"

import faiss
import numpy as np

# FIX: Use the embedding model singleton to save RAM
# Import from services/embeddings.py instead of loading SentenceTransformer directly
from .embeddings import get_embedding_model, get_embedding_dim

from routing.config import DB_PATH, logger


_kb_build_lock = Lock()
_kb_faiss_index = None
_kb_faiss_ids = []
_kb_tenant_indexes: Dict[str, tuple] = {}


def _resolve_tenant_id(tenant_id=None) -> str:
    """The tenant a search/write is scoped to. Falls back to the registered
    default tenant or active test tenant so legacy callers keep serving the
    intended data instead of leaking across tenants."""
    tid = str(tenant_id or "").strip()
    if not tid:
        try:
            from shared.tenancy.resolver import resolve_default_tenant
            tid = str(resolve_default_tenant())
        except Exception:
            tid = "default"
    try:
        from shared.tenancy import store
        rec = store.get_tenant(tid)
        if rec:
            return str(rec.get("slug") or rec.get("id"))
    except Exception:
        pass
    return tid


def _safe_json_loads(value: str, default=None):
    if default is None:
        default = []
    try:
        return json.loads(value) if value else default
    except (json.JSONDecodeError, TypeError):
        return default


CHUNK_SIZE = 800
CHUNK_OVERLAP = 150


def chunk_text(text: str) -> List[str]:
    if not text or not text.strip():
        return []
    chunks = []
    start = 0
    text_len = len(text)
    max_iterations = text_len // max(1, CHUNK_SIZE - CHUNK_OVERLAP) + 100
    iteration = 0
    while start < text_len and iteration < max_iterations:
        iteration += 1
        end = min(start + CHUNK_SIZE, text_len)
        if end < text_len:
            break_points = ["\n\n", "\n", ". ", "; ", ", "]
            search_area = text[start + int(CHUNK_SIZE * 0.3):end]
            found = False
            for bp in break_points:
                idx = search_area.rfind(bp)
                if idx != -1:
                    end = start + int(CHUNK_SIZE * 0.3) + idx + len(bp)
                    found = True
                    break
            if not found:
                pass
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        next_start = end - CHUNK_OVERLAP
        if next_start <= start:
            next_start = start + max(1, CHUNK_SIZE - CHUNK_OVERLAP)
        start = next_start
        if start >= text_len:
            break
    return chunks


def chunk_by_sections(text: str) -> List[str]:
    if not text or not text.strip():
        return []

    header_pattern = re.compile(r"^#{1,3}\s+.+", re.MULTILINE)
    if header_pattern.search(text):
        sections = header_pattern.split(text)
        sections = [s.strip() for s in sections if s.strip()]
    else:
        sections = re.split(r"\n{2,}", text)
        sections = [s.strip() for s in sections if s.strip()]

    chunks = []
    for section in sections:
        if len(section) <= CHUNK_SIZE:
            chunks.append(section)
        else:
            chunks.extend(chunk_text(section))
    return chunks


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\t", " ")
    text = re.sub(r" +", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def build_chunks_for_kb():
    """Chunk all KB docs and store chunks_json in the database."""
    try:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT id, content, chunks_json FROM knowledge_base ORDER BY id"
        ).fetchall()
        for row in rows:
            if row["chunks_json"]:
                continue
            content = row["content"] or ""
            if not content.strip():
                continue
            chunks = chunk_by_sections(content)
            if not chunks:
                chunks = chunk_text(content)
            conn.execute(
                "UPDATE knowledge_base SET chunks_json = ? WHERE id = ?",
                (json.dumps(chunks), row["id"]),
            )
        conn.commit()
        conn.close()
        logger.info(f"Chunks built for {len([r for r in rows if not r['chunks_json']])} KB docs")
    except Exception as e:
        logger.error(f"Failed to build chunks for KB: {e}")


def extract_text_from_file(file_path: str) -> str:
    """Extract text content from a file (PDF, DOCX, TXT, CSV, XLSX)."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt" or ext == ".csv":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    if ext == ".pdf":
        try:
            import pdfplumber
            texts = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        texts.append(page_text)
                    tables = page.extract_tables()
                    for table in tables:
                        if not table or len(table) < 2:
                            continue
                        rows = []
                        for row in table:
                            cleaned = [c.strip() if c else "" for c in row]
                            rows.append(" | ".join(cleaned))
                        header = rows[0]
                        sep = " | ".join(["---"] * len(table[0]))
                        table_text = "\n".join([header, sep] + rows[1:])
                        texts.append(table_text)
            combined = "\n\n".join(texts)
            if combined.strip():
                return combined
        except Exception as e:
            logger.error(f"PDF table extract via pdfplumber failed: {e}")
        try:
            from pdfminer.high_level import extract_text
            return extract_text(file_path, laparams=None) or ""
        except Exception as e:
            logger.error(f"PDF extract failed: {e}")
            try:
                import fitz
                doc = fitz.open(file_path)
                return "\n".join(page.get_text() or "" for page in doc)
            except Exception as e2:
                logger.error(f"PDF extract fitz fallback failed: {e2}")
                return ""

    if ext in (".docx", ".doc"):
        try:
            from docx import Document
            doc = Document(file_path)
            return "\n".join(para.text for para in doc.paragraphs)
        except Exception as e:
            logger.error(f"DOCX extract failed: {e}")
            return ""

    if ext in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True)
            texts = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    texts.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(texts)
        except Exception as e:
            logger.error(f"XLSX extract failed: {e}")
            return ""

    return ""


import faiss
import numpy as np


def _get_embeddings(texts: List[str]) -> np.ndarray:
    model = get_embedding_model()
    embeddings = model.encode(texts, normalize_embeddings=True)
    return embeddings.astype("float32")


def _normalize_chunks(chunks) -> List[str]:
    texts = []
    for c in chunks:
        if isinstance(c, str):
            if c.strip():
                texts.append(c.strip())
        elif isinstance(c, dict):
            text = c.get("content") or c.get("text") or ""
            text = text.strip()
            if text:
                texts.append(text)
    return texts


def add_document_to_kb(filename: str, text: str, category: str = "", media_url: str = None, media_type: str = None, tenant_id: str = None) -> Dict:
    """Add an uploaded document to the knowledge base with chunks and embeddings."""
    text = clean_text(text)
    chunks = chunk_by_sections(text)
    if not chunks:
        chunks = chunk_text(text)
    if not chunks:
        return {"success": False, "error": "No text extracted"}

    embeddings = _get_embeddings(chunks)
    embedding_blob = embeddings.tobytes()

    conn = sqlite3.connect(DB_PATH)
    try:
        cur = conn.execute(
            "INSERT INTO knowledge_base (title, category, content, chunks_json, tags, source, embedding_blob, media_url, media_type, tenant_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (filename, category, text, json.dumps(chunks), json.dumps([]), filename, embedding_blob, media_url, media_type, _resolve_tenant_id(tenant_id)),
        )
        doc_id = cur.lastrowid
        conn.commit()
    finally:
        conn.close()

    return {"success": True, "chunks": len(chunks), "doc_id": doc_id}


def rebuild_faiss_index():
    """Rebuild the FAISS index from all KB documents, partitioning by tenant."""
    global _kb_faiss_index, _kb_faiss_ids, _kb_tenant_indexes
    try:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT id, tenant_id, chunks_json, embedding_blob FROM knowledge_base"
        ).fetchall()
        conn.close()

        missing = [r for r in rows if r["embedding_blob"] is None]
        if missing:
            logger.info("rebuild_faiss_index: generating embeddings for %d docs", len(missing))
            generate_missing_embeddings()

        rows_with_emb = [r for r in rows if r["embedding_blob"] is not None]
        if not rows_with_emb:
            logger.info("rebuild_faiss_index: no documents with embeddings found")
            _kb_faiss_index = None
            _kb_faiss_ids = []
            _kb_tenant_indexes = {}
            return

        dim = get_embedding_dim()
        all_embeddings = []
        all_ids = []
        tenant_data = {}

        for row in rows_with_emb:
            chunks = _safe_json_loads(row["chunks_json"], [])
            if not chunks:
                continue
            emb = np.frombuffer(row["embedding_blob"], dtype="float32").reshape(len(chunks), dim)
            all_embeddings.append(emb)
            tid = row["tenant_id"] or "default"
            chunk_records = [(row["id"], i, tid) for i in range(len(chunks))]
            all_ids.extend(chunk_records)

            if tid not in tenant_data:
                tenant_data[tid] = {"embs": [], "ids": []}
            tenant_data[tid]["embs"].append(emb)
            tenant_data[tid]["ids"].extend(chunk_records)

        if not all_embeddings:
            logger.info("rebuild_faiss_index: no chunks with embeddings found")
            return

        # Build global index
        all_embeddings_matrix = np.vstack(all_embeddings).astype("float32")
        index = faiss.IndexFlatIP(dim)
        faiss.normalize_L2(all_embeddings_matrix)
        index.add(all_embeddings_matrix)
        _kb_faiss_index = index
        _kb_faiss_ids = all_ids

        # Build per-tenant indexes
        new_tenant_indexes = {}
        for tid, t_dict in tenant_data.items():
            if t_dict["embs"]:
                t_matrix = np.vstack(t_dict["embs"]).astype("float32")
                t_idx = faiss.IndexFlatIP(dim)
                faiss.normalize_L2(t_matrix)
                t_idx.add(t_matrix)
                new_tenant_indexes[tid] = (t_idx, t_dict["ids"])
        _kb_tenant_indexes = new_tenant_indexes

        logger.info(
            "rebuild_faiss_index: %d chunks indexed across %d tenants from %d docs",
            len(all_ids), len(_kb_tenant_indexes), len(rows_with_emb)
        )
    except Exception as e:
        logger.error(f"rebuild_faiss_index failed: {e}")


def generate_missing_embeddings():
    """Generate embeddings for KB documents that are missing them."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT id, chunks_json FROM knowledge_base WHERE embedding_blob IS NULL"
        ).fetchall()
        if not rows:
            conn.close()
            return
        logger.info("generate_missing_embeddings: %d docs need embeddings", len(rows))
        updated = 0
        for row in rows:
            chunks = _safe_json_loads(row["chunks_json"], [])
            if not chunks:
                continue
            texts = _normalize_chunks(chunks)
            if not texts:
                continue
            embeddings = _get_embeddings(texts)
            embedding_blob = embeddings.tobytes()
            conn.execute(
                "UPDATE knowledge_base SET embedding_blob = ? WHERE id = ?",
                (embedding_blob, row["id"]),
            )
            updated += 1
        conn.commit()
        logger.info("generate_missing_embeddings: updated %d docs", updated)
    finally:
        conn.close()


def search_kb_faiss(query: str, limit: int = 5, tenant_id: str = None) -> List[Dict]:
    """Search KB using FAISS vector similarity, scoped strictly to one tenant."""
    global _kb_faiss_index, _kb_faiss_ids, _kb_tenant_indexes
    tid = _resolve_tenant_id(tenant_id)
    model = get_embedding_model()
    query_emb = model.encode([query], normalize_embeddings=True).astype("float32")
    faiss.normalize_L2(query_emb)

    # 1. Prefer dedicated per-tenant index (guarantees zero cross-tenant starvation)
    t_index = None
    t_ids = None
    if _kb_tenant_indexes:
        for candidate_key in (tid, str(tid), int(tid) if str(tid).isdigit() else None):
            if candidate_key is not None and candidate_key in _kb_tenant_indexes:
                t_index, t_ids = _kb_tenant_indexes[candidate_key]
                break
        if not t_index:
            try:
                from shared.tenancy import store
                rec = store.get_tenant(tid)
                if rec:
                    for alt_key in (rec.get("id"), str(rec.get("id")), rec.get("slug")):
                        if alt_key in _kb_tenant_indexes:
                            t_index, t_ids = _kb_tenant_indexes[alt_key]
                            break
            except Exception:
                pass

    if t_index is not None and t_index.ntotal > 0:
        fetch_k = min(limit, t_index.ntotal)
        scores, indices = t_index.search(query_emb, fetch_k)
        results = []
        for score, idx in zip(scores[0], indices[0]):
            if idx < 0 or idx >= len(t_ids):
                continue
            doc_id, chunk_idx, _ = t_ids[idx]
            results.append({
                "doc_id": doc_id,
                "chunk_index": chunk_idx,
                "score": float(score),
            })
        return results

    # 2. Fall back to global index with tenant filter
    if _kb_faiss_index is None or not _kb_faiss_ids:
        return []

    fetch_k = min(limit * 4, _kb_faiss_index.ntotal)
    scores, indices = _kb_faiss_index.search(query_emb, fetch_k)
    results = []
    allowed_tids = {str(tid), int(tid) if str(tid).isdigit() else tid}
    try:
        from shared.tenancy import store
        rec = store.get_tenant(tid)
        if rec:
            allowed_tids.add(rec.get("id"))
            allowed_tids.add(str(rec.get("id")))
            allowed_tids.add(rec.get("slug"))
    except Exception:
        pass

    for score, idx in zip(scores[0], indices[0]):
        if idx < 0 or idx >= len(_kb_faiss_ids):
            continue
        doc_id, chunk_idx, doc_tenant = _kb_faiss_ids[idx]
        if doc_tenant not in allowed_tids and str(doc_tenant) not in allowed_tids:
            continue
        results.append({
            "doc_id": doc_id,
            "chunk_index": chunk_idx,
            "score": float(score),
        })
        if len(results) >= limit:
            break
    return results


# ── Product FAISS index ─────────────────────────────────────────────────

_product_faiss_index = None
_product_faiss_ids = []
_product_tenant_indexes: Dict[str, tuple] = {}


def _product_searchable_text(row) -> str:
    """Build a single searchable string from a product row for embedding."""
    parts = [
        row.get("name") or "",
        row.get("description") or "",
        row.get("short_description") or "",
        row.get("category") or "",
    ]
    ingredients = row.get("ingredients")
    if isinstance(ingredients, str):
        try:
            ingredients = json.loads(ingredients)
        except (json.JSONDecodeError, TypeError):
            pass
    if isinstance(ingredients, list):
        parts.extend(ingredients)
    elif ingredients:
        parts.append(str(ingredients))
    nutritional = row.get("nutritional_facts") or ""
    if nutritional:
        parts.append(nutritional)
    return " ".join(p for p in parts if p).strip()


def generate_product_embeddings():
    """Generate embeddings for products that are missing them."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT id, name, description, short_description, category, ingredients, nutritional_facts "
            "FROM products WHERE is_active = 1 AND embedding IS NULL"
        ).fetchall()
        if not rows:
            return
        logger.info("generate_product_embeddings: %d products need embeddings", len(rows))
        model = get_embedding_model()
        updated = 0
        for row in rows:
            text = _product_searchable_text(dict(row))
            if not text:
                continue
            emb = model.encode([text], normalize_embeddings=True).astype("float32")
            conn.execute(
                "UPDATE products SET embedding = ? WHERE id = ?",
                (emb.tobytes(), row["id"]),
            )
            updated += 1
        conn.commit()
        logger.info("generate_product_embeddings: updated %d products", updated)
    finally:
        conn.close()


def rebuild_product_faiss_index():
    """Rebuild the in-memory FAISS index for products, partitioned per tenant."""
    global _product_faiss_index, _product_faiss_ids, _product_tenant_indexes
    try:
        generate_product_embeddings()

        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT id, tenant_id, embedding FROM products WHERE is_active = 1 AND embedding IS NOT NULL"
        ).fetchall()
        conn.close()

        if not rows:
            logger.info("rebuild_product_faiss_index: no products with embeddings")
            _product_faiss_index = None
            _product_faiss_ids = []
            _product_tenant_indexes = {}
            return

        dim = get_embedding_dim()
        all_embeddings = []
        all_ids = []
        tenant_data = {}

        for row in rows:
            emb = np.frombuffer(row["embedding"], dtype="float32").reshape(1, dim)
            all_embeddings.append(emb)
            tid = row["tenant_id"] or "default"
            rec = (row["id"], tid)
            all_ids.append(rec)
            if tid not in tenant_data:
                tenant_data[tid] = {"embs": [], "ids": []}
            tenant_data[tid]["embs"].append(emb)
            tenant_data[tid]["ids"].append(rec)

        all_embeddings_matrix = np.vstack(all_embeddings).astype("float32")
        index = faiss.IndexFlatIP(dim)
        faiss.normalize_L2(all_embeddings_matrix)
        index.add(all_embeddings_matrix)

        _product_faiss_index = index
        _product_faiss_ids = all_ids

        new_tenant_indexes = {}
        for tid, t_dict in tenant_data.items():
            if t_dict["embs"]:
                t_matrix = np.vstack(t_dict["embs"]).astype("float32")
                t_idx = faiss.IndexFlatIP(dim)
                faiss.normalize_L2(t_matrix)
                t_idx.add(t_matrix)
                new_tenant_indexes[tid] = (t_idx, t_dict["ids"])
        _product_tenant_indexes = new_tenant_indexes

        logger.info(
            "rebuild_product_faiss_index: %d products indexed across %d tenants",
            len(all_ids), len(_product_tenant_indexes)
        )
    except Exception as e:
        logger.error(f"rebuild_product_faiss_index failed: {e}")


def search_products_faiss(query: str, limit: int = 5, tenant_id: str = None) -> List[Dict]:
    """Search products using FAISS vector similarity, scoped strictly to one tenant."""
    global _product_faiss_index, _product_faiss_ids, _product_tenant_indexes
    tid = _resolve_tenant_id(tenant_id)
    model = get_embedding_model()
    query_emb = model.encode([query], normalize_embeddings=True).astype("float32")
    faiss.normalize_L2(query_emb)

    # 1. Prefer dedicated per-tenant index
    if _product_tenant_indexes and tid in _product_tenant_indexes:
        t_index, t_ids = _product_tenant_indexes[tid]
        if t_index.ntotal > 0:
            fetch_k = min(limit, t_index.ntotal)
            scores, indices = t_index.search(query_emb, fetch_k)
            results = []
            for score, idx in zip(scores[0], indices[0]):
                if idx < 0 or idx >= len(t_ids):
                    continue
                product_id, _ = t_ids[idx]
                results.append({
                    "product_id": product_id,
                    "score": float(score),
                })
            return results

    # 2. Fall back to global index with tenant filter
    if _product_faiss_index is None or not _product_faiss_ids:
        return []

    scores, indices = _product_faiss_index.search(query_emb, min(limit * 4, _product_faiss_index.ntotal))
    results = []
    for score, idx in zip(scores[0], indices[0]):
        if idx < 0 or idx >= len(_product_faiss_ids):
            continue
        product_id, product_tenant = _product_faiss_ids[idx]
        if product_tenant != tid:
            continue
        results.append({
            "product_id": product_id,
            "score": float(score),
        })
        if len(results) >= limit:
            break
    return results[:limit]
