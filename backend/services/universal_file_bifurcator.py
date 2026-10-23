*** Begin File: universal_file_bifurcator.py ***
"""Universal file ingestion & image extraction.

This module replaces the PDF‑only logic with a dispatcher that can handle
- PDF
- Microsoft Word (.docx)
- Plain text (.txt)
- HTML (.html)
- Stand‑alone image files (jpg, png, gif, ...)

It returns a tuple ``(text, images)`` where ``text`` is a plain‑string that will be fed
to the existing LLM bifurcation prompt and ``images`` is a list of dicts compatible
with the original ``_extract_page_assets`` output:

```
{
    "url": "/uploaded_files/product_images/<filename>",
    "width": <int>,   # optional – 0 if unknown
    "height": <int>,  # optional – 0 if unknown
}
```

The function is deliberately lightweight and does **not** write any database records –
that part remains in ``universal_pdf_bifurcator.bifurcate_and_ingest_pdf``.
"""

import io
import mimetypes
import uuid
from enum import Enum
from pathlib import Path
from typing import List, Tuple

import fitz  # PyMuPDF – PDF handling
import docx  # python‑docx – DOCX handling
from bs4 import BeautifulSoup  # HTML parsing
from PIL import Image

# ---------------------------------------------------------------------------
# Directory where product images are stored (same as PDF engine).
# The original PDF engine defines these globals – we replicate the path logic.
# This file lives in ``backend/services``, so ``BASE_DIR`` points two levels up.
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent.parent
PRODUCT_IMAGES_DIR = BASE_DIR / "uploaded_files" / "product_images"
PRODUCT_IMAGES_DIR.mkdir(parents=True, exist_ok=True)


class FileKind(Enum):
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    HTML = "html"
    IMAGE = "image"
    UNKNOWN = "unknown"


def _detect_kind(filename: str, data: bytes) -> FileKind:
    """Return a :class:`FileKind` based on mime‑type or magic bytes.
    ``filename`` is used for extension‑based hints; ``data`` is examined for
    magic numbers when the mime guess is inconclusive.
    """
    mime, _ = mimetypes.guess_type(filename.lower())
    if mime:
        if mime == "application/pdf":
            return FileKind.PDF
        if mime in {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}:
            return FileKind.DOCX
        if mime.startswith("image/"):
            return FileKind.IMAGE
        if mime == "text/plain":
            return FileKind.TXT
        if mime == "text/html":
            return FileKind.HTML
    # Fallback – check magic bytes
    if data[:4] == b"%PDF":
        return FileKind.PDF
    if data[:2] == b"PK":  # zip container – likely DOCX
        try:
            docx.Document(io.BytesIO(data))
            return FileKind.DOCX
        except Exception:
            pass
    return FileKind.UNKNOWN


def _save_image(blob: bytes, tenant_id: str, suffix: str) -> dict:
    """Save *blob* to ``uploaded_files/product_images`` and return the dict expected
    by the rest of the pipeline.
    """
    filename = f"{tenant_id}_{uuid.uuid4().hex[:6]}{suffix}"
    path = PRODUCT_IMAGES_DIR / filename
    path.write_bytes(blob)
    # Try to obtain dimensions – not essential, but nice to have
    try:
        with Image.open(io.BytesIO(blob)) as img:
            w, h = img.size
    except Exception:
        w, h = 0, 0
    return {"url": f"/uploaded_files/product_images/{filename}", "width": w, "height": h}


def _extract_pdf(data: bytes, tenant_id: str) -> Tuple[str, List[dict]]:
    doc = fitz.open(stream=data, filetype="pdf")
    full_text = []
    images: List[dict] = []
    for page in doc:
        full_text.append(page.get_text("text"))
        # Re‑use the PDF image extraction logic from the original engine
        img_list = page.get_images(full=True)
        for img_idx, img_info in enumerate(img_list):
            xref = img_info[0]
            base = doc.extract_image(xref)
            if not base:
                continue
            raw = base.get("image")
            ext = (base.get("ext") or "jpg").lower()
            if ext == "jpeg":
                ext = "jpg"
            images.append(_save_image(raw, tenant_id, f".{ext}"))
    return "\n".join(full_text), images


def _extract_docx(data: bytes, tenant_id: str) -> Tuple[str, List[dict]]:
    doc = docx.Document(io.BytesIO(data))
    # Text extraction – concatenate paragraph texts
    text = "\n".join(p.text for p in doc.paragraphs if p.text)
    images: List[dict] = []
    # Inline pictures are stored in the underlying zip package; iterate relations.
    for i, rel in enumerate(doc.part.rels.values()):
        if "image" in rel.reltype:
            img_bytes = rel.target_part.blob
            suffix = Path(rel.target_ref).suffix or ".png"
            images.append(_save_image(img_bytes, tenant_id, suffix))
    return text, images


def _extract_text(data: bytes) -> Tuple[str, List[dict]]:
    return data.decode(errors="ignore"), []


def _extract_html(data: bytes) -> Tuple[str, List[dict]]:
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    return text, []


def _extract_image(data: bytes, filename: str, tenant_id: str) -> Tuple[str, List[dict]]:
    # Stand‑alone image – no text, just return the saved image dict.
    suffix = Path(filename).suffix or ".png"
    return "", [_save_image(data, tenant_id, suffix)]


def process(file_bytes: bytes, filename: str, tenant_id: str) -> Tuple[str, List[dict]]:
    """Detect the file type, extract plain text **and** any embedded images.
    Returns ``(text, images)`` where ``images`` matches the schema used by the
    original PDF pipeline.
    """
    kind = _detect_kind(filename, file_bytes)
    if kind == FileKind.PDF:
        return _extract_pdf(file_bytes, tenant_id)
    if kind == FileKind.DOCX:
        return _extract_docx(file_bytes, tenant_id)
    if kind == FileKind.TXT:
        return _extract_text(file_bytes)
    if kind == FileKind.HTML:
        return _extract_html(file_bytes)
    if kind == FileKind.IMAGE:
        return _extract_image(file_bytes, filename, tenant_id)
    raise ValueError(f"Unsupported file type for '{filename}'.")

*** End File: universal_file_bifurcator.py ***